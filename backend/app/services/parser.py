import PyPDF2
from pptx import Presentation
import io
from typing import Optional

def extract_text_from_pdf(file_content: bytes) -> str:
    """Extract text from PDF file"""
    try:
        pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_content))
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text() + "\n"
        return text.strip()
    except Exception as e:
        raise ValueError(f"Error parsing PDF: {str(e)}")

def extract_text_from_pptx(file_content: bytes) -> str:
    """Extract text from PowerPoint file"""
    try:
        prs = Presentation(io.BytesIO(file_content))
        text = ""
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
        return text.strip()
    except Exception as e:
        raise ValueError(f"Error parsing PowerPoint: {str(e)}")

def extract_text_from_txt(file_content: bytes) -> str:
    """Extract text from plain text file"""
    try:
        return file_content.decode('utf-8').strip()
    except Exception as e:
        raise ValueError(f"Error parsing text file: {str(e)}")

def extract_text_from_file(filename: str, file_content: bytes) -> str:
    """Route to appropriate parser based on file extension"""
    filename_lower = filename.lower()
    
    if filename_lower.endswith('.pdf'):
        return extract_text_from_pdf(file_content)
    elif filename_lower.endswith(('.ppt', '.pptx')):
        return extract_text_from_pptx(file_content)
    elif filename_lower.endswith('.txt'):
        return extract_text_from_txt(file_content)
    else:
        raise ValueError(f"Unsupported file type: {filename}")
