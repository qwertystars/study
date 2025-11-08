"""File parsing utilities for different file types"""
import os
import tempfile
from typing import Optional
import requests
import yt_dlp
from PyPDF2 import PdfReader
from pptx import Presentation


class FileParser:
    """Utility class for parsing different file types"""

    @staticmethod
    def parse_pdf(file_path: str) -> str:
        """Extract text from a PDF file"""
        try:
            reader = PdfReader(file_path)
            text = ""
            for page in reader.pages:
                text += page.extract_text() + "\n"
            return text.strip()
        except Exception as e:
            raise ValueError(f"Error parsing PDF: {str(e)}")

    @staticmethod
    def parse_pptx(file_path: str) -> str:
        """Extract text from a PowerPoint file"""
        try:
            prs = Presentation(file_path)
            text = ""
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
            return text.strip()
        except Exception as e:
            raise ValueError(f"Error parsing PowerPoint: {str(e)}")

    @staticmethod
    def parse_text(file_path: str) -> str:
        """Read a plain text file"""
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read().strip()
        except Exception as e:
            raise ValueError(f"Error reading text file: {str(e)}")

    @staticmethod
    def parse_youtube(url: str) -> str:
        """Extract transcript from a YouTube video"""
        try:
            ydl_opts = {
                "writesubtitles": True,
                "writeautomaticsub": True,
                "subtitleslangs": ["en"],
                "skip_download": True,
                "outtmpl": tempfile.mktemp(),
            }

            with yt_dlp.YoutubeDL(ydl_opts) as ydl:
                info = ydl.extract_info(url, download=False)

                # Try to get subtitles
                if "subtitles" in info and "en" in info["subtitles"]:
                    subtitle_url = info["subtitles"]["en"][0]["url"]
                    response = requests.get(subtitle_url)
                    return response.text
                elif "automatic_captions" in info and "en" in info["automatic_captions"]:
                    subtitle_url = info["automatic_captions"]["en"][0]["url"]
                    response = requests.get(subtitle_url)
                    return response.text
                else:
                    # Fallback to description
                    return info.get("description", "No transcript available")

        except Exception as e:
            raise ValueError(f"Error parsing YouTube video: {str(e)}")

    @staticmethod
    def parse_file(file_path: str, file_type: Optional[str] = None) -> str:
        """
        Parse a file based on its type

        Args:
            file_path: Path to the file
            file_type: Optional file type override (pdf, pptx, txt, youtube)

        Returns:
            Extracted text content
        """
        if not file_type:
            # Determine file type from extension
            _, ext = os.path.splitext(file_path)
            ext = ext.lower()

            if ext == ".pdf":
                file_type = "pdf"
            elif ext in [".pptx", ".ppt"]:
                file_type = "pptx"
            elif ext in [".txt", ".md"]:
                file_type = "txt"
            else:
                file_type = "txt"  # Default to text

        # Parse based on type
        if file_type == "pdf":
            return FileParser.parse_pdf(file_path)
        elif file_type == "pptx":
            return FileParser.parse_pptx(file_path)
        elif file_type == "youtube":
            return FileParser.parse_youtube(file_path)
        else:
            return FileParser.parse_text(file_path)
